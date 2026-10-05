"""Turn the LaTeX placeholder paragraphs into native, editable Word equation objects.

LaTeX -> MathML (latex2mathml) -> OMML (Office's own MML2OMML.XSL) -> injected into the paragraph.
The journal requires editable equations; this produces the same object type as Word's equation editor.
"""
import re, sys, copy
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import latex2mathml.converter as l2m

XSL = r'C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'
transform = etree.XSLT(etree.parse(XSL))
FONT = 'Times New Roman'

def latex_to_omml(latex):
    mml = l2m.convert(latex)
    omml = transform(etree.fromstring(mml.encode()))
    return copy.deepcopy(omml.getroot())

def fix(path):
    doc = Document(path)
    n = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        m = re.match(r'^\((\d+)\)\s*(.+)$', txt, re.S)
        if not m: continue
        num, latex = m.group(1), m.group(2).strip()
        if not re.search(r'[\^_{}]', latex): continue          # not an equation
        try:
            om = latex_to_omml(latex)
        except Exception as e:
            print('  skip eq', num, type(e).__name__, str(e)[:60]); continue
        for r in list(p.runs): r._element.getparent().remove(r._element)
        run = p.add_run(f'({num})\t')
        run.font.name = FONT; run.font.size = Pt(12)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        p._p.append(om)
        n += 1
    doc.save(path)
    print(f'{path.split("/")[-1]}: {n} equations converted')

if __name__ == '__main__':
    for f in sys.argv[1:]: fix(f)
