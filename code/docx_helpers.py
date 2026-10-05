from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
FONT='Microsoft JhengHei'
def setfont(run,size=None,bold=None,color=None):
    run.font.name=FONT; run._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)
def new_doc():
    d=Document()
    sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    for a in ('left_margin','right_margin'): setattr(sec,a,Cm(2.2))
    sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.2)
    for st in ['Normal','Heading 1','Heading 2','Heading 3','Title','Caption','List Bullet','List Number']:
        s=d.styles[st]; s.font.name=FONT; s.element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    d.styles['Normal'].font.size=Pt(10.5)
    d.styles['Normal'].paragraph_format.line_spacing=1.35
    d.styles['Normal'].paragraph_format.space_after=Pt(4)
    for st,sz,col in [('Heading 1',16,'1F3A5F'),('Heading 2',13,'1F3A5F'),('Heading 3',11.5,'2A5A8A'),('Title',22,'1F3A5F')]:
        d.styles[st].font.size=Pt(sz); d.styles[st].font.color.rgb=RGBColor.from_string(col)
    # page number footer
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(); f1=OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'),'begin')
    it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text='PAGE'
    f2=OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'end')
    r._r.append(f1); r._r.append(it); r._r.append(f2); setfont(r,9)
    return d
def para(d,text,bold_prefix=None,size=None,style=None,align=None):
    p=d.add_paragraph(style=style)
    if bold_prefix: setfont(p.add_run(bold_prefix),size,True)
    # simple **bold** markup
    parts=text.split('**')
    for i,t in enumerate(parts):
        if t: setfont(p.add_run(t),size,i%2==1)
    if align=='center': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
def bullet(d,text): return para(d,text,style='List Bullet')
def numbered(d,text): return para(d,text,style='List Number')
def shade(cell,hexcol):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcol); tcPr.append(sh)
def table(d,header,rows,widths=None,size=9,highlight_rows=()):
    t=d.add_table(rows=1+len(rows),cols=len(header)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(header):
        c=t.rows[0].cells[j]; c.text=''; setfont(c.paragraphs[0].add_run(str(h)),size,True); shade(c,'DCE8F5')
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.rows[i+1].cells[j]; c.text=''
            txt=str(v); b=txt.startswith('**'); txt=txt.strip('*')
            setfont(c.paragraphs[0].add_run(txt),size,b)
            c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            if i in highlight_rows: shade(c,'FFF4E0')
    if widths:
        t.autofit=False
        tblPr=t._tbl.tblPr
        lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
        tw=OxmlElement('w:tblW'); tw.set(qn('w:w'),str(int(sum(widths)*567))); tw.set(qn('w:type'),'dxa')
        for e in tblPr.findall(qn('w:tblW')): tblPr.remove(e)
        tblPr.append(tw)
        grid=t._tbl.tblGrid
        for j,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[j]*567)))
        for row in t.rows:
            for j,w in enumerate(widths): row.cells[j].width=Cm(w)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.1; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def caption(d,text):
    p=d.add_paragraph(); setfont(p.add_run(text),9,False,'52514E'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; return p
def figure(d,path,cap,width=16):
    d.add_picture(path,width=Cm(width)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; caption(d,cap)
