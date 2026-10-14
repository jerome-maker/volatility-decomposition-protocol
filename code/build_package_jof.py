"""Assemble the submission package for Journal of Forecasting (Wiley Research Exchange).

Differences from the DSFE package: review is single-blind, so the manuscript carries the author
details and there is no anonymised copy or separate title page; figures go up as separate files
rather than only embedded; and the reference list is APA.
"""
import os, re, shutil, zipfile, datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import md2docx_jof as M

D = r'J:/我的雲端硬碟/2026工作區/指導學生/柳宇軒/'
PKG = D + '投稿包_JOF/'
CODE = D + '資訊階梯分析_程式與圖表/'
FONT = 'Times New Roman'
MAIL = 'jerome.chou@ntust.edu.tw'
REPO = 'https://github.com/jerome-maker/volatility-decomposition-protocol'
FIG_ORDER = [fn for _, fn, _ in M.FIGMAP]
DOC_SCRIPTS = {'md2docx.py', 'md2docx_jof.py', 'eqfix.py', 'make_anon.py', 'build_package.py',
               'build_package_jof.py', 'build_repo.py', 'submission_check.py', 'check_jof.py',
               'refcheck.py', 'refcheck2.py', 'refcross.py', 'docx_helpers.py', 'build_report2.py',
               'dsfe_style.py', 'jof_style.py', 'apa_refs.py'}
LOCALPATH = 'J:/我的雲端硬碟/2026工作區/指導學生/柳宇軒/'


def setfont(run, size=12, bold=False):
    run.font.name = FONT; run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size); run.bold = bold


def cover_letter(path):
    d = Document()
    st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for a in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(d.sections[0], a, Cm(2.5))

    def p(txt, after=10, align=None, size=12, bold=False):
        par = d.add_paragraph(); par.paragraph_format.space_after = Pt(after)
        par.paragraph_format.line_spacing = Pt(16)
        if align == 'right': par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == 'justify': par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        setfont(par.add_run(txt), size, bold); return par

    p(datetime.date.today().strftime('%d %B %Y'), align='right', after=18)
    p('The Editor-in-Chief', after=2); p('Journal of Forecasting', after=18)
    p('Dear Editor,', after=12)
    p('I am submitting the manuscript "Decomposing deep learning gains in volatility forecasting: Information, '
      'architecture and reproducibility" for consideration as a research paper.', align='justify')
    p('Comparisons between deep learning and econometric models in financial forecasting rarely hold the '
      'information set fixed. When a neural network is given intraday measures, option-implied indices and '
      'macro-financial series that its econometric benchmark does not see, the reported gain confounds the value '
      'of the inputs with the value of the model. The paper proposes a protocol that separates them: models are '
      'arranged in a ladder whose rungs differ in the information they may use, the deep learning model and its '
      'linear counterpart receive identical inputs, and the out-of-sample gain then splits into three components '
      'that sum to it exactly — high-frequency information, exogenous forward-looking information, and the '
      'nonlinear architecture.', align='justify')
    p('Applied to five-day realized volatility of the SPDR S&P 500 exchange-traded fund, with a second '
      'application to the Invesco QQQ fund, only the exogenous component improves the QLIKE loss significantly. '
      'Two results bear directly on how such comparisons should be conducted. The statistical and economic '
      'rankings diverge: the model with the lowest loss ranks third on risk-adjusted return under volatility '
      'targeting, and the models are indistinguishable in value-at-risk backtesting. And the architecture '
      'component is of the same order as the variation caused by the number of random seeds and by the version '
      'of the deep learning library, so the paper closes with four minimum standards for comparisons of this '
      'kind. A hold-out year, reserved before the design was frozen and used once, reverses the sign of that '
      'component and is reported as it stands.', align='justify')
    p('The work speaks to the journal\'s emphasis on methodological and computational contributions to '
      'forecasting, and on the relationship between forecasting systems and the decisions they serve: the '
      'economic evaluation in Section 6 is where the statistical ranking and the portfolio ranking part company. '
      'It also connects to recent work on whether machine learning beats linear benchmarks once the fitting '
      'scheme is held fixed (Chassot & Audrino, 2026; Branco et al., 2024), to which this paper is the '
      'complement: it fixes the fitting scheme, varies the information set, and decomposes the difference.',
      align='justify')
    p('All code, the daily derived series, the forecast series for every model and seed, and a locked '
      f'specification of the software environment are available at {REPO}, so that every table and figure can be '
      'recomputed. The manuscript has not been published elsewhere and is not under consideration by another '
      'journal. I have no competing interests to declare, the research received no external funding, and I am '
      'not suggesting reviewers.', align='justify')
    p('Generative AI tools (Claude, Anthropic) were used for code drafting, data processing, drafting of the '
      'manuscript text, language editing and document preparation, and in discussing analytical choices. All '
      'research decisions, analyses and interpretations were made and verified by me, and I take full '
      'responsibility for the content.', align='justify')
    p('Thank you for considering this submission.', after=18, align='justify')
    p('Yours sincerely,', after=24)
    p('Jerome Chih-Lung Chou', after=2)
    p('Graduate Institute of Finance, National Taiwan University of Science and Technology', size=11, after=2)
    p(MAIL, size=11)
    d.save(path)


def zip_items(path, items):
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for arc, src in items:
            z.write(src, arc)


def main():
    KEEP = {'00_README.md', '01_manuscript.docx', '02_manuscript.pdf', '03_figures.zip',
            '04_supplementary.zip', '05_cover_letter.docx', 'figures', 'desktop.ini', '_tmp'}
    os.makedirs(PKG, exist_ok=True); os.makedirs(PKG + 'figures', exist_ok=True)
    for f in os.listdir(PKG):
        if f in KEEP: continue
        q = PKG + f
        try:
            shutil.rmtree(q, ignore_errors=True) if os.path.isdir(q) else os.remove(q)
            print('removed stale', f)
        except OSError as e:
            print('COULD NOT REMOVE (close it?):', f, e.strerror)

    shutil.copy(D + 'Chou_2026_JOF_manuscript.docx', PKG + '01_manuscript.docx')
    shutil.copy(D + 'Chou_2026_JOF_manuscript.pdf', PKG + '02_manuscript.pdf')
    # figures both loose (Wiley uploads one file per figure) and zipped
    for i, fn in enumerate(FIG_ORDER, 1):
        shutil.copy(CODE + 'fig/' + fn, PKG + f'figures/Figure{i}.png')
    zip_items(PKG + '03_figures.zip', [(f'Figure{i}.png', CODE + 'fig/' + fn)
                                       for i, fn in enumerate(FIG_ORDER, 1)])
    tmp = PKG + '_tmp/'; os.makedirs(tmp, exist_ok=True)
    supp = [('ENVIRONMENT.md', CODE + 'ENVIRONMENT.md')]
    for f in sorted(os.listdir(CODE)):
        if not f.endswith('.py') or f in DOC_SCRIPTS: continue
        src = open(CODE + f, encoding='utf-8').read().replace(LOCALPATH, './')
        open(tmp + f, 'w', encoding='utf-8').write(src)
        supp.append(('code/' + f, tmp + f))
    for f in sorted(os.listdir(CODE)):
        if f.endswith(('.csv', '.json')): supp.append(('results/' + f, CODE + f))
    for f in sorted(os.listdir(CODE + 'data')):
        if f.endswith(('.csv', '.json')): supp.append(('data/' + f, CODE + 'data/' + f))
    for f in sorted(os.listdir(CODE + 'predictions')):
        if f.endswith('.csv'): supp.append(('data/predictions/' + f, CODE + 'predictions/' + f))
    zip_items(PKG + '04_supplementary.zip', supp)
    shutil.rmtree(tmp, ignore_errors=True)
    cover_letter(PKG + '05_cover_letter.docx')

    doc = Document(PKG + '01_manuscript.docx')
    body = '\n'.join(p.text for p in doc.paragraphs)
    total = sum(os.path.getsize(os.path.join(r, f)) for r, _, fs in os.walk(PKG) for f in fs)
    readme = f"""# Journal of Forecasting submission package

Title: {M.TITLE}
Running title: {M.RUNNING} ({len(M.RUNNING)} characters, limit 40)
Author: Jerome Chih-Lung Chou, Graduate Institute of Finance, National Taiwan University of Science and Technology
Corresponding e-mail: {MAIL}
Prepared: {datetime.date.today()}   Review model: single-blind (so the manuscript carries author details)

## What to upload in Research Exchange

| Order | File | Purpose |
|---|---|---|
| 1 | `01_manuscript.docx` | Main document: title page, abstract, keywords, body, references, tables, figures, figure legends |
| 2 | `02_manuscript.pdf` | PDF of the same document, for editors who prefer it |
| 3 | `figures/Figure1–9.png` or `03_figures.zip` | Figures as separate files, 300 dpi, each ≥ 2299 px wide |
| 4 | `04_supplementary.zip` | Code, derived data, forecast series and the locked environment |
| 5 | `05_cover_letter.docx` | Cover letter |

The journal charges no submission fee; an article publishing charge applies only if open access
is chosen after acceptance.

## Manuscript statistics

- Pages 29   Words (body, excluding tables) {len(re.findall(r"[A-Za-z][A-Za-z'’-]+", body)):,}
- Tables {len(doc.tables)}   Figures {len(FIG_ORDER)}   Equation objects {len(doc.part.element.xpath('.//m:oMath'))}
- Abstract 277 words, 6 keywords (Wiley's limit), references 47 in APA style
- Declarations in the manuscript: data availability statement, acknowledgments (no funding),
  conflict of interest, use of AI tools

Total package {total / 1e6:.1f} MB.

## What changed from the earlier version of this manuscript

- Title and abstract lead with the decomposition and the reproducibility findings rather than with
  information-set parity.
- References rebuilt in APA style, the style this journal uses, with authors, volume, issue and pages
  taken from Crossref; five author lists were corrected in the process.
- In-text citations follow APA: "et al." from three authors, an ampersand inside parentheses.
- Eleven tables that were never called out in the running text now are.
- One reference that nothing in the text cited was removed.
- No anonymised copy or separate title page: review here is single-blind.

## Before uploading

1. Research Exchange asks for an ORCID identifier; it is not in any file here.
2. Figures must be uploaded as separate files with the legend entered beneath each one; the legends are
   also listed at the end of the manuscript, as the guidelines require.
3. Confirm the corresponding address and affiliation on the title page before submitting.
"""
    open(PKG + '00_README.md', 'w', encoding='utf-8').write(readme)
    print(readme)
    for f in sorted(os.listdir(PKG)):
        q = PKG + f
        if os.path.isfile(q): print(f'  {f:28s} {os.path.getsize(q)/1e6:7.2f} MB')
    print(f'  total {total/1e6:.1f} MB at {PKG}')


if __name__ == '__main__':
    main()
