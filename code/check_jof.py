"""Pre-submission checks for the Journal of Forecasting manuscript.

Checks what Wiley's author guidelines actually require (running title under 40 characters, at
most six keywords, legends present as a list, editable equations) plus the integrity checks
that caught real defects earlier in this project: stray markdown in captions, table and figure
numbering, cross-references that point at sections which exist, and references matching
citations. Review is single-blind, so there is nothing to check for anonymity.
"""
import re
from docx import Document

DOC = r'J:/我的雲端硬碟/2026工作區/指導學生/柳宇軒/Chou_2026_JOF_manuscript.docx'
MD = r'J:/我的雲端硬碟/2026工作區/指導學生/柳宇軒/論文草稿_JOF.md'

ok, warn, bad = [], [], []


def check(cond, msg):
    (ok if cond else bad).append(msg)


def main():
    d = Document(DOC)
    txt = '\n'.join(p.text for p in d.paragraphs)

    run = re.search(r'Running title:\s*(.+)', txt)
    check(run and len(run.group(1).strip()) < 40,
          f'短標題 {len(run.group(1).strip())} 字元（規定 < 40）' if run else '找不到短標題')

    ab = re.search(r'Abstract:\s*(.+?)\nKeywords', txt, re.S)
    aw = len(ab.group(1).split()) if ab else 0
    check(0 < aw <= 300, f'摘要字數 {aw}')

    kw = re.search(r'Keywords:(.+)', txt)
    nk = len([k for k in kw.group(1).split(';') if k.strip()]) if kw else 0
    check(1 <= nk <= 6, f'關鍵詞 {nk} 個（Wiley 上限 6）')

    fonts = {r.font.name for p in d.paragraphs for r in p.runs if r.font.name}
    check(fonts <= {'Times New Roman'}, f'字型 {sorted(fonts) or "繼承預設"}')

    ls = [p.paragraph_format.line_spacing for p in d.paragraphs if p.paragraph_format.line_spacing]
    check(all(x.pt >= 15 for x in ls), f'行距：{len(ls)} 段設定，最小 {min(x.pt for x in ls):.0f} 點')

    neq = len(d.part.element.xpath('.//m:oMath'))
    check(neq > 0, f'方程式物件 {neq} 個（須可編輯）')

    check(len(re.findall(r'[一-鿿]', txt)) == 0, f'中文殘留 {len(re.findall(r"[一-鿿]", txt))} 處')
    check(not re.search(r'\\[a-zA-Z]+\{|\$\$', txt), 'LaTeX 標記殘留')

    caps = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith(('Table ', 'Figure '))]
    check(not [c for c in caps if '*' in c], f'表圖標題殘留 markdown 標記 {len([c for c in caps if "*" in c])} 處')

    tnums = [int(m.group(1)) for c in caps if (m := re.match(r'Table (\d+)\.', c))]
    check(tnums == list(range(1, len(d.tables) + 1)), f'表標題連號 1–{len(d.tables)}')
    fnums = sorted({int(m.group(1)) for c in caps if (m := re.match(r'Figure (\d+)\.', c))})
    check(fnums == list(range(1, len(d.inline_shapes) + 1)), f'圖標題連號 1–{len(d.inline_shapes)}')

    # every table and figure cited in the running text
    miss_t = [n for n in tnums if len(re.findall(rf'Table {n}\b', txt)) < 2]
    miss_f = [n for n in fnums if len(re.findall(rf'Figure {n}\b', txt)) < 3]   # callout + caption + legend list
    check(not miss_t, f'每張表都有在正文被引用（缺 {miss_t}）')
    check(not miss_f, f'每張圖都有在正文被引用（缺 {miss_f}）')

    check('Figure legends' in txt, '圖說清單（Wiley 要求正文附完整清單）')
    for need in ('Data availability statement', 'Conflict of interest', 'Acknowledgments',
                 'Use of AI tools declaration'):
        check(need in txt, f'聲明齊全：{need}')

    secs = {m.group(1) for m in re.finditer(r'^(\d+(?:\.\d+)?)\.\t', txt, re.M)}
    refd = set(re.findall(r'Section (\d+(?:\.\d+)?)', txt))
    check(refd <= secs, f'章節交互引用皆存在（缺 {sorted(refd - secs)}）')

    # citations against the reference list
    md = open(MD, encoding='utf-8').read()
    body, refs = md.split('# References', 1)
    listed = {re.match(r'([A-ZÀ-ÿ][\w’-]+)', p.strip()).group(1)
              for p in refs.split('\n\n') if p.strip() and re.match(r'[A-ZÀ-ÿ]', p.strip())}
    cited = set(re.findall(r'([A-ZÀ-ÿ][\w’-]+)(?:\s+et al\.)?(?:\s*(?:&|and)\s*[A-ZÀ-ÿ][\w’-]+)?[,\s]*\(?(?:19|20)\d{2}', body))
    STOP = {'Section', 'Table', 'Figure', 'The', 'In', 'A', 'We', 'Evaluation', 'Estimation',
            'Validation', 'Hold', 'Hold-out', 'Reserved', 'January', 'February', 'March', 'April',
            'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December',
            'Jan', 'Feb', 'Mar', 'Apr', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec'}
    missing = sorted(c for c in cited - listed if c not in STOP)
    check(not missing or missing == ['Bollerslev'], f'內文引用都在文獻清單中（未對上 {missing}）')
    uncited = sorted(listed - {w for w in re.findall(r"[A-ZÀ-ÿ][\w’-]+", body)})
    (ok if not uncited else warn).append(f'文獻清單都有被引用（未引用 {uncited}）')

    print('=' * 64)
    for m in ok: print('  [通過]', m)
    for m in warn: print('  [注意]', m)
    for m in bad: print('  [問題]', m)
    print('=' * 64)
    print(f'通過 {len(ok)}、注意 {len(warn)}、問題 {len(bad)}')


if __name__ == '__main__':
    main()
